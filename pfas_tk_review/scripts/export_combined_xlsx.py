#!/usr/bin/env python3
"""Write the four combined datasets into one workbook, one sheet per table.

The CSVs in db/combined/ are the authoritative form; this is a convenience
export for opening alongside the original PFAS_TK.xlsx.

Run:  python3 scripts/export_combined_xlsx.py
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "db", "combined")
DEST = os.path.join(HERE, "..", "db", "combined", "PFAS_TK_combined.xlsx")

SHEETS = [
    ("tk_parameters.csv", "TK parameters"),
    ("binding.csv", "Protein binding"),
    ("transporters.csv", "Transporters"),
    ("regulatory.csv", "Regulatory"),
]
HEAD_FILL = PatternFill("solid", fgColor="2A78D6")


def main():
    wb = Workbook()
    wb.remove(wb.active)

    readme = wb.create_sheet("README")
    for i, line in enumerate([
        ["PFAS toxicokinetics review - combined datasets"],
        [],
        ["Built by scripts/build_combined_datasets.py from the per-pass tables in"],
        ["db/ and the per-paper extractions in db/primary_2026/."],
        [],
        ["Every row carries a `provenance` column naming the file it came from, and"],
        ["a study / pmid_or_doi / source_table trio naming the primary source."],
        ["Nothing here is quoted from memory."],
        [],
        ["Sheet", "Rows", "What it holds"],
    ], 1):
        readme.append(line)
    readme["A1"].font = Font(bold=True, size=14)
    readme["A10"].font = readme["B10"].font = readme["C10"].font = Font(bold=True)

    for fn, title in SHEETS:
        path = os.path.join(SRC, fn)
        with open(path) as fh:
            rows = list(csv.reader(fh))
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        for c in range(1, len(rows[0]) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEAD_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            width = max((len(str(rw[c - 1])) for rw in rows[:400]
                         if c - 1 < len(rw)), default=10)
            ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 10), 46)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        readme.append([title, len(rows) - 1, fn])

    wb.save(DEST)
    print("wrote", os.path.relpath(DEST, os.path.join(HERE, "..")))


if __name__ == "__main__":
    main()
