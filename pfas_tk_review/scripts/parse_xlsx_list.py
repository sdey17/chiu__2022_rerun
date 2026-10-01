#!/usr/bin/env python3
"""Parse the user-supplied PFAS_TK.xlsx citation list, split each citation into
fields, and classify every entry by the kind of evidence it can contribute to a
toxicokinetics review.

The point of the classification is to make explicit which of the 81 entries can
supply a toxicokinetic parameter (half-life, Vd, clearance, binding constant,
transporter Km) and which are hazard/mechanism studies that cannot.
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(
    "/root/.claude/uploads/602cc481-1215-5cea-847e-a573be17cdd9/411ec763-PFAS_TK.xlsx"
)
OUT = Path(__file__).resolve().parent.parent / "db" / "xlsx_source_list.csv"

# Mojibake from the spreadsheet's encoding, mapped back to the intended letters.
MOJIBAKE = {
    "√©": "e", "√≠": "i", "√°": "a", "√∂": "o", "√§": "a", "√•": "a",
    "√º": "u", "√®": "e", "Œ∫": "kappa", "Œ±": "alpha", "Œ≤": "beta",
    "√ñ": "O", "√¶": "ae",
}

# Ordered most specific first: the first pattern that matches wins, so a paper
# that measures both binding and toxicity lands in the binding bucket.
CATEGORIES = [
    ("in_vitro_protein_binding", r"serum albumin|albumin|globulin|binding of|binding affinit|"
                                 r"binding to|lipid binding|transthyretin|binds?\b"),
    ("transporter_or_enzyme", r"organic anion transport|oatp|\boat\b|transporter|uptake of|"
                              r"carboxylesterase|cytochrome p450|cyp2|biotransformation enzyme|"
                              r"renal transport|renal excretion|phase i and ii"),
    ("in_vivo_tk", r"pharmacokinetic|toxicokinetic|half-life|elimination|distribution of|"
                   r"bioconcentration|bioaccumulation|chemical burden|internal concentration|"
                   r"dermal penetration|tissue distribution|excretion after|urinary metabolite"),
    ("receptor_binding", r"estrogen receptor|ppar|thyroid hormone t3 receptor|"
                         r"g protein-coupled estrogen receptor|nuclear receptor"),
    ("membrane_physchem", r"bilayer|phase transition|membrane injury|cellular localization|"
                          r"plasma membrane"),
]


def demojibake(s: str) -> str:
    for bad, good in MOJIBAKE.items():
        s = s.replace(bad, good)
    return s


def split_citation(raw: str):
    """Return (authors, year, title, journal) from 'A,B,C YYYY. Title. Journal.'"""
    s = demojibake(raw).strip()
    m = re.search(r"\s(19|20)\d{2}\.\s", s)
    if not m:
        return s, "", "", ""
    year = s[m.start() + 1 : m.start() + 5]
    authors = s[: m.start()].strip().rstrip(",")
    rest = s[m.end() :].strip()
    # The journal is the final short segment; titles can contain periods, so take
    # the last segment only when it looks like an abbreviated journal name.
    parts = [p.strip() for p in rest.rstrip(".").split(". ")]
    if len(parts) >= 2 and len(parts[-1]) < 60:
        title, journal = ". ".join(parts[:-1]), parts[-1]
    else:
        title, journal = rest.rstrip("."), ""
    return authors, year, title.rstrip("."), journal


def classify(title: str, journal: str) -> str:
    hay = f"{title} {journal}".lower()
    for name, pat in CATEGORIES:
        if re.search(pat, hay):
            return name
    return "mechanistic_hazard"


def species(title: str) -> str:
    hay = title.lower()
    found = []
    for pat, label in [
        (r"\brats?\b|sprague|wistar|f344", "rat"),
        (r"\bmice\b|\bmouse\b|c57bl|cd-1|balb", "mouse"),
        (r"human|hela|hepg2|caco-2|\bhsa\b", "human"),
        (r"zebrafish|danio", "zebrafish"),
        (r"bovine|\bbsa\b", "bovine"),
        (r"macaca|monkey|primate", "primate"),
        (r"chicken|frog|seal|trout|elegans", "other_animal"),
    ]:
        if re.search(pat, hay):
            found.append(label)
    return "|".join(found) if found else "unspecified"


def chemicals(title: str) -> str:
    hay = title.lower()
    hits = []
    for pat, label in [
        (r"perfluorooctanoic|\bpfoa\b|perfluorooctanoate|ammonium perfluorooctanoate|"
         r"perfluorooctane acid", "PFOA"),
        (r"perfluorooctanesulfon|perfluorooctane sulfon|\bpfos\b|perfluorooctanesulfonic", "PFOS"),
        (r"perfluorononano|\bpfna\b", "PFNA"),
        (r"perfluorobutane sulfon|perfluorobutanesulfon|\bpfbs\b", "PFBS"),
        (r"perfluorobutyrate|perfluorobutanoic|\bpfba\b", "PFBA"),
        (r"perfluorohexane sulfon|perfluorohexanesulfon|\bpfhxs\b", "PFHxS"),
        (r"perfluorohexano|\bpfhxa\b", "PFHxA"),
        (r"perfluorodecano|perfluoro-n-decanoic|\bpfda\b", "PFDA"),
        (r"perfluoroalkyl|perfluorinated|polyfluoroalkyl|\bpfas\b|perfluorocarboxylic|"
         r"perfluoro fatty", "multiple_PFAS"),
    ]:
        if re.search(pat, hay):
            hits.append(label)
    # Drop the generic label when a specific compound was also named.
    if len(hits) > 1 and "multiple_PFAS" in hits:
        hits.remove("multiple_PFAS")
    return "|".join(hits) if hits else "unspecified"


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    raws = [str(r[0]).strip() for r in ws.iter_rows(values_only=True)
            if r[0] and str(r[0]).strip()]

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["row", "first_author", "year", "title", "journal",
                    "category", "species_in_title", "chemicals_in_title",
                    "tk_parameter_possible", "raw_citation"])
        counts = {}
        for i, raw in enumerate(raws, 1):
            authors, year, title, journal = split_citation(raw)
            cat = classify(title, journal)
            first = authors.split(",")[0].strip() if authors else ""
            # Only these three categories can yield a number a TK model consumes.
            tk = "yes" if cat in {"in_vivo_tk", "in_vitro_protein_binding",
                                  "transporter_or_enzyme"} else "no"
            counts[cat] = counts.get(cat, 0) + 1
            w.writerow([i, first, year, title, journal, cat, species(title),
                        chemicals(title), tk, demojibake(raw)])

    print(f"wrote {OUT}  ({len(raws)} rows)")
    print("\ncategory counts")
    for k, v in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"  {k:28} {v:3}")


if __name__ == "__main__":
    main()
